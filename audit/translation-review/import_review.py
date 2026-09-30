"""Apply a reviewed workbook back to the site.
Rows marked "Fix" with a "Corrected translation" replace that string in antz-onboard/lang/<code>.json;
then lang/build.py checks and rebuilds. Add --reviewed to mark the language as reviewed (removes the
machine-translation note), once the reviewer has been through every row.

  python3 audit/translation-review/import_review.py path/to/Antz-Guide-Translation-Review-de-German.xlsx [--reviewed]
Works with the .xlsx as returned, or downloaded from Google Sheets (File > Download > Microsoft Excel)."""
import json, os, re, subprocess, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
LANG = os.path.join(ROOT, '..', '..', 'antz-onboard', 'lang')


def set_value(d, key, value):
    part = key.split('.', 1)
    if part[0] == 'ui':
        d['ui'][part[1]] = value; return
    if part[0] == 'track':
        tid, f = part[1].rsplit('.', 1); d['tracks'][tid][f] = value; return
    mid, rest = part[1].split('.', 1)
    m = d['modules'][mid]
    if rest.startswith('shot.'):
        m.setdefault('shots', {})[rest[5:]] = value; return
    f = re.match(r'f(\d+)\.(\w+)$', rest)
    if f:
        m['features'][int(f.group(1)) - 1][f.group(2)] = value; return
    m[rest] = value


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    code = re.search(r'Review-(\w\w)-', os.path.basename(path))
    if not code:
        sys.exit('Cannot tell the language from the file name (expected ...Review-<code>-...)')
    code = code.group(1)
    ws = load_workbook(path, data_only=True)['Strings']
    head = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(head)}
    k, st, fix = col['Key (do not edit)'], col['Status'], col['Corrected translation']
    lp = os.path.join(LANG, f'{code}.json')
    d = json.load(open(lp, encoding='utf-8'))
    counts, applied, skipped = {'OK': 0, 'Fix': 0, 'Unsure': 0, '': 0}, 0, []
    for row in ws.iter_rows(min_row=3, values_only=True):
        key = row[k]
        if not key or key == 'EXAMPLE':
            continue
        status = (row[st] or '').strip()
        counts[status if status in counts else ''] += 1
        new = (row[fix] or '').strip() if row[fix] else ''
        if status == 'Fix' and new:
            try:
                set_value(d, key, new); applied += 1
            except (KeyError, IndexError):
                skipped.append(key)
    if '--reviewed' in sys.argv:
        d.setdefault('meta', {})['reviewed'] = True
    json.dump(d, open(lp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'{code}: {applied} corrections applied; {counts["OK"]} OK, {counts["Fix"]} fix, {counts["Unsure"]} unsure, {counts[""]} not reviewed')
    if skipped:
        print('Keys not found (left unchanged):', ', '.join(skipped))
    subprocess.run([sys.executable, os.path.join(LANG, 'build.py')], check=False)


if __name__ == '__main__':
    main()
